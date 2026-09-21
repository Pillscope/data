import json
from pathlib import Path
from openpyxl import load_workbook

SRC = Path.home() / "data/eu/medicines_human.xlsx"
OUT = Path.home() / "data/eu/processed"
OUT.mkdir(parents=True, exist_ok=True)

HEADER_ROW_INDEX = 8  # 0-based; real headers start here
NUM_COLUMNS = 39       # everything after this is blank padding

wb = load_workbook(SRC, read_only=True, data_only=True)
ws = wb.active
rows = ws.iter_rows(values_only=True)

for _ in range(HEADER_ROW_INDEX):
    next(rows)
header = list(next(rows))[:NUM_COLUMNS]

index = []
count = 0
for row in rows:
if not row or not row[1]:  # column 1 = Name of medicine
    continue
if row[0] != "Human":  # column 0 = Category, skip veterinary
    continue
    row = list(row)[:NUM_COLUMNS]
    record = {header[i]: row[i] for i in range(len(header))}

    product_number = record.get("EMA product number") or f"row{count}"
    slug = str(product_number).replace("/", "_").replace(" ", "_")

    with open(OUT / f"{slug}.json", "w") as f:
        json.dump(record, f, indent=2, default=str)

    index.append({
        "id": slug,
        "name": record.get("Name of medicine", ""),
        "active_substance": record.get("Active substance", ""),
        "status": record.get("Medicine status", ""),
    })
    count += 1

with open(OUT / "_index.json", "w") as f:
    json.dump(index, f, indent=2)

print(f"Processed {count} EU medicines into {OUT}")
